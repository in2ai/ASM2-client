import inspect
import logging
import re
import textwrap
from abc import ABC, abstractmethod
from io import BytesIO, StringIO
import csv

from docx import Document as DocxFile
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    Paragraph as RLParagraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table as RLTable,
    TableStyle,
)

from src.generation.model import Document


logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Input sanitising, shared by every renderer
#
# Renderers are fed model output, so no cell can be assumed to be a str, no row
# can be assumed to match the heading count, and no heading level can be assumed
# to be in range. ReportLab rejects ragged rows outright and str.join() raises on
# non-str cells, so this runs before anything else.
# ---------------------------------------------------------------------------

_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def _text(value) -> str:
    if value is None:
        return ""
    if not isinstance(value, str):
        try:
            value = str(value)
        except Exception:
            return ""
    return _CONTROL_CHARS.sub("", value)


def _level(level) -> int:
    try:
        return max(1, min(6, int(level)))
    except Exception:
        return 2


def _items(items) -> list[str]:
    if items is None:
        return []
    if isinstance(items, str):
        return [_text(items)]
    try:
        return [_text(i) for i in items]
    except TypeError:
        return [_text(items)]


def _table_data(headings, rows) -> tuple[list[str], list[list[str]]]:
    """Rectangular, all-str table data."""
    headings = _items(headings)
    try:
        rows = [_items(r) for r in (rows or [])]
    except TypeError:
        rows = []
    ncols = max([len(headings)] + [len(r) for r in rows] + [1])
    headings = headings + [""] * (ncols - len(headings))
    rows = [r + [""] * (ncols - len(r)) for r in rows]
    return headings, rows


def _wrap_lines(text: str, width: int) -> list[str]:
    """Width-aware wrap that preserves indentation and hard-breaks long tokens."""
    width = max(int(width), 8)
    out: list[str] = []
    for raw in text.expandtabs(4).splitlines() or [""]:
        if not raw.strip():
            out.append("")
            continue
        indent = raw[: len(raw) - len(raw.lstrip())][: width // 2]
        body = raw.lstrip()
        wrapped = textwrap.wrap(
            body,
            max(width - len(indent), 8),
            break_long_words=True,
            break_on_hyphens=False,
        ) or [""]
        out.extend(indent + line for line in wrapped)
    return out


class DocumentRenderer(ABC):
    # Document-level metadata, for formats that can carry it. Set per render()
    # rather than in reset(), which subclasses override without calling up.
    document_title: str = ""

    def render(self, document: Document) -> bytes:
        self.reset()
        self.document_title = document.title
        document.render(self)
        return self.finish()

    def reset(self) -> None:
        pass

    @abstractmethod
    def finish(self) -> bytes: ...

    @abstractmethod
    def heading(self, text: str, level: int) -> None: ...

    @abstractmethod
    def paragraph(self, text: str) -> None: ...

    @abstractmethod
    def item_list(self, items: list[str], ordered: bool) -> None: ...

    @abstractmethod
    def table(self, headings: list[str], rows: list[list[str]]) -> None: ...

    @abstractmethod
    def code_block(self, code: str, language: str | None) -> None: ...


# ---------------------------------------------------------------------------
# Plain text
# ---------------------------------------------------------------------------

class TxtRenderer(DocumentRenderer):
    def reset(self) -> None:
        self.blocks: list[str] = []

    def heading(self, text, level):
        text = _text(text)
        self.blocks.append(text.upper() if _level(level) == 1 else text)

    def paragraph(self, text):
        self.blocks.append(_text(text))

    def item_list(self, items, ordered):
        items = _items(items)
        lines = [f"{i}. {item}" if ordered else f"- {item}" for i, item in enumerate(items, 1)]
        self.blocks.append("\n".join(lines))

    def table(self, headings, rows):
        headings, rows = _table_data(headings, rows)
        lines = [" | ".join(headings)] + [" | ".join(row) for row in rows]
        self.blocks.append("\n".join(lines))

    def code_block(self, code, language):
        self.blocks.append(_text(code))

    def finish(self) -> bytes:
        return "\n\n".join(self.blocks).encode("utf-8")


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------

class MarkdownRenderer(DocumentRenderer):
    def reset(self) -> None:
        self.blocks: list[str] = []

    def heading(self, text, level):
        self.blocks.append(f"{'#' * _level(level)} {_text(text)}")

    def paragraph(self, text):
        self.blocks.append(_text(text))

    def item_list(self, items, ordered):
        items = _items(items)
        lines = [f"{i}. {item}" if ordered else f"- {item}" for i, item in enumerate(items, 1)]
        self.blocks.append("\n".join(lines))

    def table(self, headings, rows):
        headings, rows = _table_data(headings, rows)

        # A pipe in a cell would silently add a column.
        def esc(cell: str) -> str:
            return cell.replace("|", "\\|").replace("\n", " ")

        header = "| " + " | ".join(esc(h) for h in headings) + " |"
        sep = "| " + " | ".join("---" for _ in headings) + " |"
        body = ["| " + " | ".join(esc(c) for c in row) + " |" for row in rows]
        self.blocks.append("\n".join([header, sep, *body]))

    def code_block(self, code, language):
        code = _text(code)
        # Longer fence than any run of backticks inside the code.
        runs = [len(m) for m in re.findall(r"`+", code)] or [0]
        fence = "`" * max(3, max(runs) + 1)
        self.blocks.append(f"{fence}{_text(language)}\n{code}\n{fence}")

    def finish(self) -> bytes:
        return "\n\n".join(self.blocks).encode("utf-8")


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

_ACCENT_HEX = "#2b3a55"
_ACCENT = colors.HexColor(_ACCENT_HEX)
_ROW_SHADE = colors.HexColor("#f4f4f6")
_CODE_BG = colors.HexColor("#f5f5f5")
_LINE = colors.HexColor("#cccccc")

_MARGIN = 0.9 * inch
_FRAME_PAD = 6.0  # SimpleDocTemplate's Frame default, on each side.

# The usable box, matching what the frame reports: LETTER minus margins minus
# frame padding, i.e. 470.4 x 650.4pt.
_AVAIL_W = LETTER[0] - 2 * _MARGIN - 2 * _FRAME_PAD
_AVAIL_H = LETTER[1] - 2 * _MARGIN - 2 * _FRAME_PAD

# A row taller than this reads badly even when it fits, and a row taller than the
# frame cannot be laid out at all, so past this point tables get linearised.
_ROW_LIMIT = 0.6 * _AVAIL_H

_TABLE_KW = set(inspect.signature(RLTable.__init__).parameters)


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _link_sub(m: re.Match) -> str:
    href = m.group(2).replace('"', "%22")
    return f'<link href="{href}" color="{_ACCENT_HEX}">{m.group(1)}</link>'


def _inline_markup(text: str) -> str:
    """Turn the inline Markdown our schema allows into ReportLab's mini-markup."""
    text = _escape(text)
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"\*(.+?)\*", r"<i>\1</i>", text)
    text = re.sub(r"`(.+?)`", r'<font face="Courier">\1</font>', text)
    text = re.sub(r"\[(.+?)\]\((.+?)\)", _link_sub, text)
    return text


def _longest_word(text: str, font: str, size: float) -> float:
    return max((stringWidth(w, font, size) for w in text.split()), default=0.0)


def _col_widths(headings, rows, font, size, pad, avail) -> list[float]:
    """Weight columns by their natural width.

    With colWidths=None ReportLab sizes each column from its *minimum* content
    width (roughly the longest word), which is what collapsed the 8-column table
    into a 699pt row. Explicit widths are the actual fix.
    """
    ncols = len(headings)
    gutter = 2 * pad + 1  # cell padding plus the grid line
    space = max(avail - ncols * gutter, ncols * 4.0)

    natural, floors = [], []
    for c in range(ncols):
        cells = [headings[c]] + [r[c] for r in rows]
        natural.append(max((stringWidth(x, font, size) for x in cells), default=0.0) or 1.0)
        floors.append(min(max((_longest_word(x, font, size) for x in cells), default=0.0), space / ncols))

    # sqrt damping: a verbose column earns more room, but not all of it.
    weights = [n ** 0.5 for n in natural]
    total = sum(weights) or 1.0
    widths = [space * w / total for w in weights]

    # Lift anything under its unbreakable minimum, taking the slack from columns
    # that have room to spare.
    for _ in range(ncols):
        deficit = sum(max(0.0, f - w) for f, w in zip(floors, widths))
        if deficit <= 0.5:
            break
        donors = [i for i in range(ncols) if widths[i] > floors[i]]
        pool = sum(widths[i] - floors[i] for i in donors)
        if pool <= 0.5:
            break
        take = min(deficit, pool)
        for i in donors:
            widths[i] -= take * (widths[i] - floors[i]) / pool
        widths = [max(w, f) for w, f in zip(widths, floors)]

    scale = space / (sum(widths) or 1.0)
    return [w * scale + gutter for w in widths]


def _row_heights(table) -> list[float] | None:
    """Ask ReportLab what the rows measure, so we can degrade before it raises.

    None means the measurement itself failed, which is a reliable sign that
    build() would fail too.
    """
    try:
        table.wrap(_AVAIL_W, _AVAIL_H)
        return [float(h or 0.0) for h in (table._rowHeights or [])]
    except Exception:
        logger.warning("Table measurement failed", exc_info=True)
        return None


class _Layout:
    """One rung of the ladder in PdfRenderer.finish()."""

    def __init__(self, name, font_scale=1.0, tables=True, boxes=True, repeat_headers=True):
        self.name = name
        self.font_scale = font_scale
        self.tables = tables            # False: every table becomes labelled blocks
        self.boxes = boxes              # False: code blocks lose their shaded box
        self.repeat_headers = repeat_headers

    def table_font(self, ncols: int) -> float:
        base = 9.0 if ncols <= 4 else 8.0 if ncols <= 6 else 7.0 if ncols <= 9 else 6.5
        return max(5.5, base * self.font_scale)

    def table_pad(self, ncols: int) -> float:
        return 8.0 if ncols <= 4 else 4.0 if ncols <= 6 else 2.5


_LAYOUTS = (
    _Layout("full"),
    _Layout("compact", font_scale=0.85, repeat_headers=False),
    _Layout("linear", font_scale=0.85, tables=False, boxes=False, repeat_headers=False),
)


class PdfRenderer(DocumentRenderer):
    def __init__(self):
        styles = getSampleStyleSheet()
        styles.add(ParagraphStyle("DocTitle", parent=styles["Title"], textColor=_ACCENT, spaceAfter=20))
        styles.add(ParagraphStyle("SectionHeading", parent=styles["Heading2"], textColor=_ACCENT, spaceBefore=16, spaceAfter=8))
        styles.add(ParagraphStyle("Body", parent=styles["BodyText"], spaceAfter=8, leading=15))
        styles.add(ParagraphStyle("RowTitle", parent=styles["Body"], textColor=_ACCENT, spaceBefore=8, spaceAfter=2))
        styles.add(ParagraphStyle("Field", parent=styles["Body"], leftIndent=14, spaceAfter=2, leading=13))
        self.styles = styles

    # -- collection: records semantic blocks only, so it cannot raise ---------

    def reset(self) -> None:
        self.blocks: list[tuple] = []

    def heading(self, text, level):
        self.blocks.append(("heading", _text(text), _level(level)))

    def paragraph(self, text):
        self.blocks.append(("paragraph", _text(text)))

    def item_list(self, items, ordered):
        self.blocks.append(("list", _items(items), bool(ordered)))

    def table(self, headings, rows):
        headings, rows = _table_data(headings, rows)
        self.blocks.append(("table", headings, rows))

    def code_block(self, code, language):
        self.blocks.append(("code", _text(code), _text(language) or None))

    # -- output --------------------------------------------------------------

    def finish(self) -> bytes:
        for layout in _LAYOUTS:
            pdf = self._build(layout)
            if pdf:
                return pdf
        logger.error("Every platypus layout failed for %r; using the plain-text floor",
                     self.document_title)
        try:
            return self._plain_text_pdf()
        except Exception:
            logger.exception("Plain-text PDF floor failed for %r", self.document_title)
            # Last resort. If this raises too, ReportLab is unusable and the
            # exception should propagate: a zero-byte .pdf is worse than a 500.
            return self._notice_pdf()

    def _build(self, layout: _Layout) -> bytes | None:
        try:
            story = self._story(layout)
            buf = BytesIO()
            doc = SimpleDocTemplate(
                buf, pagesize=LETTER,
                # Without these, ReportLab stamps its "(anonymous)" placeholders, and
                # PDF viewers label the tab with that instead of the document title.
                title=self.document_title or "Document",
                author="",
                subject="",
                creator="ASM2",
                topMargin=_MARGIN, bottomMargin=_MARGIN,
                leftMargin=_MARGIN, rightMargin=_MARGIN,
            )
            doc.build(story)
            return buf.getvalue()
        except Exception:
            logger.warning("PDF layout %s failed for %r; trying the next rung",
                           layout.name, self.document_title, exc_info=True)
            return None

    def _story(self, layout: _Layout) -> list:
        story: list = []
        for block in self.blocks:
            try:
                story.extend(self._flowables(block, layout))
            except Exception:
                logger.warning("Degrading a %s block to plain text", block[0], exc_info=True)
                story.extend(self._plain_block(block))
        return story or [Spacer(1, 1)]

    def _flowables(self, block: tuple, layout: _Layout) -> list:
        kind = block[0]
        if kind == "heading":
            style = "DocTitle" if block[2] == 1 else "SectionHeading"
            return [self._para(block[1], self.styles[style])]
        if kind == "paragraph":
            return [self._para(block[1], self.styles["Body"])]
        if kind == "list":
            return self._list(block[1], block[2])
        if kind == "table":
            return self._table(block[1], block[2], layout)
        if kind == "code":
            return self._code(block[1], layout)
        return []

    def _para(self, text: str, style, prefix: str = ""):
        """prefix is trusted markup; text is escaped and marked up."""
        try:
            return RLParagraph(prefix + _inline_markup(text), style)
        except Exception:
            # Markup the paragraph parser rejected: fall back to the literal.
            return RLParagraph(prefix + _escape(text), style)

    def _list(self, items, ordered):
        if not items:
            return []
        rows = [ListItem(self._para(i, self.styles["Body"])) for i in items]
        return [
            ListFlowable(rows, bulletType="1" if ordered else "bullet", leftIndent=18),
            Spacer(1, 8),
        ]

    def _table(self, headings, rows, layout: _Layout):
        if not layout.tables or not rows:
            return self._linear_table(headings, rows)

        ncols = len(headings)
        size = layout.table_font(ncols)
        pad = layout.table_pad(ncols)
        cell = ParagraphStyle("Cell", parent=self.styles["Body"],
                              fontSize=size, leading=size * 1.28, spaceAfter=0)
        head = ParagraphStyle("Head", parent=cell,
                              fontName="Helvetica-Bold", textColor=colors.white)

        header = [self._para(h, head) for h in headings]
        body = [[self._para(c, cell) for c in row] for row in rows]

        kwargs = dict(
            colWidths=_col_widths(headings, rows, cell.fontName, size, pad, _AVAIL_W),
            hAlign="LEFT",
            splitByRow=1,
        )
        if "splitInRow" in _TABLE_KW:
            # Off deliberately. Splitting inside a row makes ReportLab re-emit the
            # repeated header mid-page and leaves the continuation fragment with a
            # blank first column, so the reader sees an unnamed row. Rows are
            # capped at _ROW_LIMIT below, well under the frame height, so no row
            # can ever be too tall to place whole and this costs nothing.
            kwargs["splitInRow"] = 0
        if "emptyTableAction" in _TABLE_KW:
            # Unreachable given the guards above, but rl_config defaults this to
            # 'error'. 'ignore' degrades to a zero Spacer; 'indicate' would draw
            # red "Table(0,0)" text into a customer-facing document.
            kwargs["emptyTableAction"] = "ignore"

        t = RLTable([header, *body], **kwargs)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), _ACCENT),
            ("GRID", (0, 0), (-1, -1), 0.5, _LINE),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, _ROW_SHADE]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), pad),
            ("RIGHTPADDING", (0, 0), (-1, -1), pad),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))

        heights = _row_heights(t)
        if heights is None or (heights and max(heights) > _ROW_LIMIT):
            return self._linear_table(headings, rows)

        # repeatRows re-emits the header on every page, so the header itself has
        # to be small enough to leave room for body rows.
        header_h = heights[0] if heights else 0.0
        t.repeatRows = 1 if (layout.repeat_headers and 0 < header_h < 0.3 * _AVAIL_H) else 0
        return [t, Spacer(1, 10)]

    def _linear_table(self, headings, rows):
        """One labelled block per row.

        Paragraphs always split across pages, so this shape cannot overflow a
        frame no matter how wide or verbose the table was.
        """
        if not rows:
            return []
        flow: list = [Spacer(1, 4)]
        for n, row in enumerate(rows, 1):
            title = row[0].strip() if row else ""
            flow.append(RLParagraph(f"<b>{_escape(title or str(n))}</b>", self.styles["RowTitle"]))
            for h, c in zip(headings, row):
                if not c.strip() or c == title:
                    continue
                label = f"<b>{_escape(h)}:</b> " if h.strip() else ""
                flow.append(self._para(c, self.styles["Field"], prefix=label))
        flow.append(Spacer(1, 8))
        return flow

    def _code(self, code: str, layout: _Layout):
        base = self.styles["Code"]
        size = max(6.0, base.fontSize * layout.font_scale)
        style = ParagraphStyle("CodeBox", parent=base, fontSize=size, leading=size * 1.22)
        pad = 10.0
        inner = _AVAIL_W - 2 * pad - 2

        # Exact for Courier, a fair approximation otherwise. Preformatted does no
        # wrapping of its own, and textwrap on a character count guessed at 90 let
        # long lines run off the page.
        char_w = (stringWidth("M" * 20, style.fontName, size) / 20) or (size * 0.6)
        lines = _wrap_lines(code, int(inner / char_w))

        if not layout.boxes:
            return [Preformatted("\n".join(lines), style), Spacer(1, 10)]

        # One box per page-sized chunk. A single-cell table is one atomic row, so
        # a code block longer than a page would otherwise be a LayoutError.
        per_page = max(1, int((_AVAIL_H - 6 * pad) / style.leading))
        flow: list = []
        for i in range(0, max(len(lines), 1), per_page):
            chunk = "\n".join(lines[i:i + per_page]) or " "
            box = RLTable([[Preformatted(chunk, style)]], colWidths=[_AVAIL_W - 2], hAlign="LEFT")
            box.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), _CODE_BG),
                ("BOX", (0, 0), (-1, -1), 0.5, _LINE),
                ("LEFTPADDING", (0, 0), (-1, -1), pad),
                ("RIGHTPADDING", (0, 0), (-1, -1), pad),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]))
            flow.extend([box, Spacer(1, 10)])
        return flow

    def _plain_block(self, block: tuple) -> list:
        """Any block as escaped Body paragraphs, with no markup left to misparse."""
        kind, body = block[0], self.styles["Body"]
        if kind == "heading":
            return [RLParagraph(_escape(block[1]), self.styles["SectionHeading"])]
        if kind == "table":
            headings, rows = block[1], block[2]
            out = []
            for row in rows:
                joined = " - ".join(f"{h}: {c}" if h.strip() else c
                                    for h, c in zip(headings, row) if c.strip())
                out.append(RLParagraph(_escape(joined), body))
            return out
        if kind == "list":
            return [RLParagraph(_escape(f"- {i}"), body) for i in block[1]]
        if kind == "code":
            return [RLParagraph(_escape(line) or " ", self.styles["Code"])
                    for line in _wrap_lines(block[1], 96)]
        return [RLParagraph(_escape(block[1]), body)]

    # -- floors: no platypus, so no layout engine and no LayoutError ---------

    def _text_lines(self, cols: int) -> list[str]:
        out: list[str] = []
        if self.document_title:
            out += [self.document_title, "=" * min(cols, len(self.document_title)), ""]
        for block in self.blocks:
            kind = block[0]
            if kind == "heading":
                out += ["", block[1].upper() if block[2] == 1 else block[1], ""]
            elif kind == "paragraph":
                out += _wrap_lines(block[1], cols) + [""]
            elif kind == "list":
                for n, item in enumerate(block[1], 1):
                    bullet = f"{n}. " if block[2] else "- "
                    wrapped = _wrap_lines(item, cols - len(bullet)) or [""]
                    out.append(bullet + wrapped[0])
                    out += [" " * len(bullet) + w for w in wrapped[1:]]
                out.append("")
            elif kind == "table":
                headings, rows = block[1], block[2]
                for row in rows:
                    for h, c in zip(headings, row):
                        if c.strip():
                            out += _wrap_lines(f"{h}: {c}" if h.strip() else c, cols)
                    out.append("")
            elif kind == "code":
                out += _wrap_lines(block[1], cols) + [""]
        return out

    def _plain_text_pdf(self) -> bytes:
        buf = BytesIO()
        c = Canvas(buf, pagesize=LETTER)
        c.setTitle(self.document_title or "Document")
        c.setCreator("ASM2")
        font, size, leading = "Courier", 9.0, 11.0
        char_w = stringWidth("M", font, size) or 5.4
        cols = max(20, int((LETTER[0] - 2 * _MARGIN) / char_w))
        y = LETTER[1] - _MARGIN
        c.setFont(font, size)
        for line in self._text_lines(cols):
            if y < _MARGIN:
                c.showPage()
                c.setFont(font, size)
                y = LETTER[1] - _MARGIN
            # Base-14 fonts only cover WinAnsi; substitute rather than fail.
            safe = line.encode("latin-1", "replace").decode("latin-1")
            c.drawString(_MARGIN, y, safe[:cols])
            y -= leading
        c.showPage()
        c.save()
        return buf.getvalue()

    def _notice_pdf(self) -> bytes:
        buf = BytesIO()
        c = Canvas(buf, pagesize=LETTER)
        c.setTitle(self.document_title or "Document")
        c.setFont("Helvetica", 11)
        c.drawString(_MARGIN, LETTER[1] - _MARGIN, "This document could not be rendered as a PDF.")
        c.showPage()
        c.save()
        return buf.getvalue()

# Same palette as the PDF, in the hex-without-hash form OOXML wants.
_D_ACCENT = _ACCENT_HEX.lstrip("#").upper()
_D_ACCENT_RGB = RGBColor.from_string(_D_ACCENT)
_D_ROW_SHADE = "F4F4F6"
_D_CODE_BG = "F5F5F5"
_D_LINE = "CCCCCC"

_BODY_FONT = "Calibri"
_MONO_FONT = "Consolas"

# Column weighting only needs relative widths, and reportlab has no Calibri
# metrics, so Helvetica stands in for it. Courier stands in for Consolas when
# wrapping code, where the two are close enough that a full line still fits.
_METRIC_FONT = "Helvetica"
_MONO_METRIC_FONT = "Courier"

_CELL_VPAD = 6.0  # matches the PDF's TOP/BOTTOMPADDING
_PAGE_MARGIN = Inches(0.9)

# Child orders from ECMA-376. Anything inserted into these parents has to land
# in this sequence or Word calls the file corrupt and refuses to open it.
_PPR_ORDER = (
    "w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
    "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd",
    "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)
_TBLPR_ORDER = (
    "w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual",
    "w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc",
    "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout",
    "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription",
)
_TCPR_ORDER = (
    "w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders",
    "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
    "w:vAlign", "w:hideMark",
)
_TRPR_ORDER = (
    "w:cnfStyle", "w:divId", "w:gridBefore", "w:gridAfter", "w:wBefore",
    "w:wAfter", "w:cantSplit", "w:trHeight", "w:tblHeader", "w:tblCellSpacing",
    "w:jc", "w:hidden",
)

# _text() already strips the C0 controls. Surrogates survive a json.loads() of
# model output and make lxml raise on assignment, so they go here.
_XML_UNSAFE = re.compile(r"[\ud800-\udfff\ufffe\uffff]")


def _xml_text(value) -> str:
    return _XML_UNSAFE.sub("", _text(value))


# -- raw OOXML plumbing -----------------------------------------------------

def _place(parent, element, order: tuple[str, ...]):
    """Insert `element` into `parent` at its schema position."""
    tag = element.tag.split("}")[-1]
    try:
        rest = order[order.index(f"w:{tag}") + 1:]
    except ValueError:
        rest = ()
    for later in rest:
        sibling = parent.find(qn(later))
        if sibling is not None:
            sibling.addprevious(element)
            return element
    parent.append(element)
    return element


def _child(parent, tag: str, order: tuple[str, ...]):
    found = parent.find(qn(tag))
    return found if found is not None else _place(parent, OxmlElement(tag), order)


def _shade(pr, order: tuple[str, ...], fill: str) -> None:
    shd = _child(pr, "w:shd", order)
    shd.set(qn("w:val"), "clear")  # "solid" fills with the *foreground* colour
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def _borders(pr, tag: str, order: tuple[str, ...], sides, color: str, eighths: int = 4):
    bdr = OxmlElement(tag)
    for side in sides:  # already in schema order
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), str(eighths))  # eighths of a point
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), color)
        bdr.append(edge)
    old = pr.find(qn(tag))
    if old is not None:
        pr.remove(old)
    return _place(pr, bdr, order)


def _table_width(table, points: float) -> None:
    tblW = _child(table._tbl.tblPr, "w:tblW", _TBLPR_ORDER)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(int(Pt(points).twips)))


def _cell_margins(table, left_right: float, top_bottom: float = _CELL_VPAD) -> None:
    mar = OxmlElement("w:tblCellMar")
    for side, pts in (("top", top_bottom), ("left", left_right),
                      ("bottom", top_bottom), ("right", left_right)):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:w"), str(int(Pt(pts).twips)))
        edge.set(qn("w:type"), "dxa")
        mar.append(edge)
    tblPr = table._tbl.tblPr
    old = tblPr.find(qn("w:tblCellMar"))
    if old is not None:
        tblPr.remove(old)
    _place(tblPr, mar, _TBLPR_ORDER)


def _repeat_header(row) -> None:
    _child(row._tr.get_or_add_trPr(), "w:tblHeader", _TRPR_ORDER).set(qn("w:val"), "true")


# -- lists ------------------------------------------------------------------

_ABSTRACT_NUM = (
    '<w:abstractNum %s w:abstractNumId="{aid}">'
    '<w:multiLevelType w:val="singleLevel"/>'
    '<w:lvl w:ilvl="0">'
    '<w:start w:val="1"/>'
    '<w:numFmt w:val="{fmt}"/>'
    '<w:lvlText w:val="{marker}"/>'
    '<w:lvlJc w:val="left"/>'
    '<w:pPr><w:ind w:left="{left}" w:hanging="{hanging}"/></w:pPr>'
    "{rpr}"
    "</w:lvl>"
    "</w:abstractNum>"
) % nsdecls("w")

_NUM = '<w:num %s w:numId="{nid}"><w:abstractNumId w:val="{aid}"/></w:num>' % nsdecls("w")

_SYMBOL_BULLET = "\uf0b7"  # Symbol font's round bullet, which is what Word writes


def _numbering_id(doc, ordered: bool) -> int | None:
    """A numbering definition of its own, per list.

    Sharing one definition would make the second ordered list in a document
    carry on counting from the first, which is the classic python-docx list bug.
    None means the caller should fall back to literal markers.
    """
    try:
        numbering = doc.part.numbering_part.element
        used_abstract = [int(e.get(qn("w:abstractNumId")) or 0)
                         for e in numbering.findall(qn("w:abstractNum"))]
        used_num = [int(e.get(qn("w:numId")) or 0)
                    for e in numbering.findall(qn("w:num"))]
        aid = max(used_abstract, default=-1) + 1
        nid = max(used_num, default=0) + 1

        abstract = parse_xml(_ABSTRACT_NUM.format(
            aid=aid,
            fmt="decimal" if ordered else "bullet",
            marker="%1." if ordered else _SYMBOL_BULLET,
            left=int(Pt(24).twips),
            hanging=int(Pt(14).twips),
            rpr="" if ordered else
                '<w:rPr><w:rFonts w:ascii="Symbol" w:hAnsi="Symbol" w:hint="default"/></w:rPr>',
        ))
        num = parse_xml(_NUM.format(nid=nid, aid=aid))

        # Every w:abstractNum has to precede every w:num.
        siblings = numbering.findall(qn("w:abstractNum"))
        if siblings:
            siblings[-1].addnext(abstract)
        else:
            numbering.insert(0, abstract)
        last_num = numbering.findall(qn("w:num"))
        if last_num:
            last_num[-1].addnext(num)
        else:
            abstract.addnext(num)
        return nid
    except Exception:
        logger.warning("No usable numbering part; falling back to literal list markers",
                       exc_info=True)
        return None


def _apply_numbering(paragraph, num_id: int) -> None:
    numPr = _child(paragraph._p.get_or_add_pPr(), "w:numPr", _PPR_ORDER)
    for tag, val in (("w:ilvl", "0"), ("w:numId", str(num_id))):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        numPr.append(el)


# -- inline markup ----------------------------------------------------------

_INLINE = re.compile(
    r"\*\*(?P<bold>.+?)\*\*"
    r"|\*(?P<italic>[^*]+?)\*"
    r"|`(?P<code>[^`]+?)`"
    r"|\[(?P<label>[^\]]*?)\]\((?P<href>[^)]*?)\)",
    re.DOTALL,
)


def _segments(text: str):
    """(text, bold, italic, mono, href) for the inline Markdown our schema allows.

    The PDF hands this to reportlab's mini-markup; Word has no markup layer, so
    the same four patterns have to become separate runs.
    """
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            yield text[pos:m.start()], False, False, False, None
        if m.group("bold") is not None:
            yield m.group("bold"), True, False, False, None
        elif m.group("italic") is not None:
            yield m.group("italic"), False, True, False, None
        elif m.group("code") is not None:
            yield m.group("code"), False, False, True, None
        else:
            href = (m.group("href") or "").strip()
            yield (m.group("label") or href), False, False, False, href or None
        pos = m.end()
    if pos < len(text):
        yield text[pos:], False, False, False, None


def _style_run(run, bold, italic, mono, size):
    run.bold = bold or None
    run.italic = italic or None
    if mono:
        run.font.name = _MONO_FONT
        run.font.size = Pt((size or 10.5) * 0.92)
    elif size:
        run.font.size = Pt(size)
    return run


# Stricter than the PDF, which passes any href straight to reportlab. A .docx
# relationship target is a live link in Word, and these are the only schemes a
# generated document has any business carrying.
_SAFE_HREF = re.compile(r"^(?:https?://|mailto:|tel:|#)", re.IGNORECASE)


def _hyperlink(paragraph, text: str, href: str, size):
    """A run wrapped in w:hyperlink, with the relationship registered."""
    if not _SAFE_HREF.match(href):
        logger.warning("Dropping link to unsupported target %r", href[:120])
        return _style_run(paragraph.add_run(text), False, False, False, size)

    run = paragraph.add_run(text)
    run.font.color.rgb = _D_ACCENT_RGB
    run.font.underline = True
    if size:
        run.font.size = Pt(size)
    try:
        r_id = paragraph.part.relate_to(href.replace('"', "%22"), RT.HYPERLINK,
                                        is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), r_id)
        link.append(run._r)  # moves the run element out of the paragraph
        paragraph._p.append(link)
    except Exception:
        # Unusable target: keep the text, lose the link.
        logger.warning("Could not link %r", href[:120], exc_info=True)
    return run


def _write(paragraph, text: str, size: float | None = None, prefix: str = ""):
    """prefix is literal text, always bold; text carries the inline markup."""
    if prefix:
        _style_run(paragraph.add_run(prefix), True, False, False, size)
    for chunk, bold, italic, mono, href in _segments(_xml_text(text)):
        if not chunk:
            continue
        if href:
            _hyperlink(paragraph, chunk, href, size)
        else:
            # python-docx turns \n into <w:br/> and \t into <w:tab/> for us.
            _style_run(paragraph.add_run(chunk), bold, italic, mono, size)
    return paragraph


# -- table sizing -----------------------------------------------------------

def _cell_font(ncols: int) -> float:
    """Mirrors _Layout.table_font so both formats shrink at the same points."""
    return 9.0 if ncols <= 4 else 8.0 if ncols <= 6 else 7.0 if ncols <= 9 else 6.5


def _cell_pad(ncols: int) -> float:
    return 8.0 if ncols <= 4 else 4.0 if ncols <= 6 else 2.5


def _estimated_heights(headings, rows, widths, size: float, pad: float) -> list[float]:
    """Roughly what Word will make of each row.

    Word does the real measuring, so this only has to be good enough to decide
    between a table and the linearised form -- the same call _row_heights()
    makes for the PDF, minus the ability to ask the layout engine.
    """
    leading = size * 1.28
    heights = []
    for row in [headings, *rows]:
        lines = 1
        for cell, width in zip(row, widths):
            usable = max(width - 2 * pad - 1, 8.0)
            lines = max(lines, sum(
                int(stringWidth(seg, _METRIC_FONT, size) // usable) + 1
                for seg in cell.split("\n")
            ))
        heights.append(lines * leading + 2 * _CELL_VPAD)
    return heights


class DocxRenderer(DocumentRenderer):
    """Word output, same block vocabulary as the other three renderers.

    Collection mirrors PdfRenderer: the visitor methods only record sanitised
    tuples, so nothing a model produces can raise before finish() gets a chance
    to degrade.
    """

    # -- collection ---------------------------------------------------------

    def reset(self) -> None:
        self.blocks: list[tuple] = []

    def heading(self, text, level):
        self.blocks.append(("heading", _xml_text(text), _level(level)))

    def paragraph(self, text):
        self.blocks.append(("paragraph", _xml_text(text)))

    def item_list(self, items, ordered):
        self.blocks.append(("list", [_xml_text(i) for i in _items(items)], bool(ordered)))

    def table(self, headings, rows):
        headings, rows = _table_data(headings, rows)
        self.blocks.append(("table",
                            [_xml_text(h) for h in headings],
                            [[_xml_text(c) for c in row] for row in rows]))

    def code_block(self, code, language):
        self.blocks.append(("code", _xml_text(code), _xml_text(language) or None))

    # -- output -------------------------------------------------------------

    def finish(self) -> bytes:
        for linear in (False, True):
            try:
                return self._build(linear)
            except Exception:
                logger.warning("Docx build (linear=%s) failed for %r; trying the next rung",
                               linear, self.document_title, exc_info=True)
        logger.error("Every docx layout failed for %r; using the plain floor",
                     self.document_title)
        # If this raises too, python-docx is unusable and the exception should
        # propagate: a zero-byte .docx is worse than a 500.
        return self._plain_docx()

    def _build(self, linear_tables: bool) -> bytes:
        doc = DocxFile()
        self._page_setup(doc)
        self._styles(doc)
        try:
            doc.core_properties.title = self.document_title or "Document"
            doc.core_properties.author = ""
        except Exception:
            logger.warning("Could not set core properties", exc_info=True)

        for block in self.blocks:
            try:
                self._emit(doc, block, linear_tables)
            except Exception:
                logger.warning("Degrading a %s block to plain text", block[0], exc_info=True)
                self._plain_block(doc, block)

        buf = BytesIO()
        doc.save(buf)
        return buf.getvalue()

    # -- document furniture -------------------------------------------------

    def _page_setup(self, doc) -> None:
        try:
            # python-docx's stock template writes a bare <w:zoom/>, which fails
            # XSD validation downstream. Word tolerates it; strict readers don't.
            zoom = doc.settings.element.find(qn("w:zoom"))
            if zoom is not None and zoom.get(qn("w:percent")) is None:
                zoom.set(qn("w:percent"), "100")
        except Exception:
            logger.warning("Could not normalise document settings", exc_info=True)

        for section in doc.sections:
            section.page_width = Inches(8.5)  # the template default is A4
            section.page_height = Inches(11)
            section.top_margin = section.bottom_margin = _PAGE_MARGIN
            section.left_margin = section.right_margin = _PAGE_MARGIN

    def _styles(self, doc) -> None:
        def get(name, base=None):
            try:
                return doc.styles[name]
            except KeyError:
                style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
                if base:
                    style.base_style = doc.styles[base]
                return style

        normal = get("Normal")
        normal.font.name = _BODY_FONT
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(8)
        normal.paragraph_format.line_spacing = 1.15

        title = get("Title")
        title.font.name = _BODY_FONT
        title.font.size = Pt(24)
        title.font.bold = True
        title.font.color.rgb = _D_ACCENT_RGB
        title.paragraph_format.space_after = Pt(20)
        # Title carries no outline level of its own, so a document whose
        # sections are all level 1 would have an empty navigation pane.
        _child(title.element.get_or_add_pPr(), "w:outlineLvl", _PPR_ORDER).set(qn("w:val"), "0")

        for level, size in ((1, 16), (2, 13), (3, 11.5), (4, 10.5), (5, 10.5)):
            try:
                h = doc.styles[f"Heading {level}"]
            except KeyError:
                continue
            h.font.name = _BODY_FONT
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.italic = level >= 4
            h.font.color.rgb = _D_ACCENT_RGB
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(8)

        cell = get("Table Cell", "Normal")
        cell.paragraph_format.space_after = Pt(0)
        cell.paragraph_format.line_spacing = 1.0

        head = get("Table Heading", "Table Cell")
        head.font.bold = True
        head.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        row_title = get("Row Title", "Normal")
        row_title.font.bold = True
        row_title.font.color.rgb = _D_ACCENT_RGB
        row_title.paragraph_format.space_before = Pt(8)
        row_title.paragraph_format.space_after = Pt(2)

        field = get("Field", "Normal")
        field.paragraph_format.left_indent = Pt(14)
        field.paragraph_format.space_after = Pt(2)

        code = get("Code Block", "Normal")
        code.font.name = _MONO_FONT
        code.font.size = Pt(8.5)
        code.paragraph_format.space_after = Pt(0)
        code.paragraph_format.line_spacing = 1.15
        # font.name only sets ascii/hAnsi; without cs, Word picks its own
        # monospace substitute for anything non-Latin in a listing.
        code.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), _MONO_FONT)

        spacer = get("Block Spacer", "Normal")
        spacer.font.size = Pt(5)
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.space_before = Pt(0)

        item = get("List Item", "Normal")
        item.paragraph_format.space_after = Pt(2)

    def _spacer(self, doc) -> None:
        doc.add_paragraph(style="Block Spacer")

    # -- blocks -------------------------------------------------------------

    def _emit(self, doc, block: tuple, linear_tables: bool) -> None:
        kind = block[0]
        if kind == "heading":
            level = block[2]
            # Level 1 is the document title, as in the PDF; deeper levels map
            # onto real Word headings so the navigation pane and any TOC work.
            style = "Title" if level == 1 else f"Heading {min(level - 1, 5)}"
            _write(doc.add_paragraph(style=self._safe_style(doc, style)), block[1])
        elif kind == "paragraph":
            _write(doc.add_paragraph(), block[1])
        elif kind == "list":
            self._list(doc, block[1], block[2])
        elif kind == "table":
            self._table(doc, block[1], block[2], linear_tables)
        elif kind == "code":
            self._code(doc, block[1])

    def _safe_style(self, doc, name: str) -> str:
        try:
            doc.styles[name]
            return name
        except KeyError:
            return "Normal"

    def _list(self, doc, items, ordered) -> None:
        if not items:
            return
        num_id = _numbering_id(doc, ordered)
        p = None
        for n, item in enumerate(items, 1):
            p = doc.add_paragraph(style=self._safe_style(doc, "List Item"))
            if num_id is None:
                # No numbering part: literal markers plus a hanging indent.
                p.paragraph_format.left_indent = Pt(24)
                p.paragraph_format.first_line_indent = Pt(-14)
                _write(p, item, prefix=f"{n}. " if ordered else "\u2022 ")
            else:
                _apply_numbering(p, num_id)
                _write(p, item)
        if p is not None:
            p.paragraph_format.space_after = Pt(10)  # the PDF's trailing Spacer

    def _table(self, doc, headings, rows, linear_tables: bool) -> None:
        if linear_tables or not rows:
            return self._linear_table(doc, headings, rows)

        ncols = len(headings)
        size = _cell_font(ncols)
        pad = _cell_pad(ncols)
        widths = _col_widths(headings, rows, _METRIC_FONT, size, pad, _AVAIL_W)

        heights = _estimated_heights(headings, rows, widths, size, pad)
        if max(heights) > _ROW_LIMIT:
            # Same call the PDF makes: past this the row reads badly, and the
            # header would be stranded a page away from its own row.
            return self._linear_table(doc, headings, rows)

        table = doc.add_table(rows=len(rows) + 1, cols=ncols)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False  # w:tblLayout fixed, so our widths are honoured
        _table_width(table, sum(widths))
        _cell_margins(table, pad)
        _borders(table._tbl.tblPr, "w:tblBorders", _TBLPR_ORDER,
                 ("top", "left", "bottom", "right", "insideH", "insideV"), _D_LINE)

        for column, width in zip(table.columns, widths):
            column.width = Pt(width)

        for c, text in enumerate(headings):
            cell = table.rows[0].cells[c]
            cell.width = Pt(widths[c])  # Word wants the width on the cell too
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            _shade(cell._tc.get_or_add_tcPr(), _TCPR_ORDER, _D_ACCENT)
            _write(self._cell_para(cell, "Table Heading"), text, size)
        if heights[0] < 0.3 * _AVAIL_H:
            _repeat_header(table.rows[0])

        for r, row in enumerate(rows, 1):
            shade = _D_ROW_SHADE if r % 2 == 0 else None
            for c, text in enumerate(row):
                cell = table.rows[r].cells[c]
                cell.width = Pt(widths[c])
                cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                if shade:
                    _shade(cell._tc.get_or_add_tcPr(), _TCPR_ORDER, shade)
                _write(self._cell_para(cell, "Table Cell"), text, size)

        self._spacer(doc)

    def _cell_para(self, cell, style_name: str):
        p = cell.paragraphs[0]
        try:
            p.style = style_name
        except KeyError:
            pass
        return p

    def _linear_table(self, doc, headings, rows) -> None:
        """One labelled block per row, for tables too wide or too tall to read."""
        if not rows:
            return
        for n, row in enumerate(rows, 1):
            title = row[0].strip() if row else ""
            _write(doc.add_paragraph(style=self._safe_style(doc, "Row Title")),
                   title or str(n))
            for heading, value in zip(headings, row):
                if not value.strip() or value == title:
                    continue
                p = doc.add_paragraph(style=self._safe_style(doc, "Field"))
                _write(p, value, prefix=f"{heading}: " if heading.strip() else "")
        self._spacer(doc)

    def _code(self, doc, code: str) -> None:
        """A one-cell table, which is how you get a box with real padding.

        Unlike the PDF this does not chunk by page: Word splits a table row
        across pages on its own, so a long listing just flows.
        """
        size = 8.5
        pad = 10.0
        inner = _AVAIL_W - 2 * pad - 2
        char_w = (stringWidth("M" * 20, _MONO_METRIC_FONT, size) / 20) or (size * 0.6)
        lines = _wrap_lines(code, int(inner / char_w)) or [""]

        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        _table_width(table, _AVAIL_W)
        _cell_margins(table, pad, 8.0)
        _borders(table._tbl.tblPr, "w:tblBorders", _TBLPR_ORDER,
                 ("top", "left", "bottom", "right"), _D_LINE)

        cell = table.rows[0].cells[0]
        cell.width = Pt(_AVAIL_W)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        _shade(cell._tc.get_or_add_tcPr(), _TCPR_ORDER, _D_CODE_BG)

        p = self._cell_para(cell, "Code Block")
        # No inline markup here: backticks and asterisks are code, not emphasis.
        p.add_run("\n".join(lines))
        self._spacer(doc)

    def _plain_block(self, doc, block: tuple) -> None:
        """Any block as literal paragraphs, with no markup left to misparse."""
        kind = block[0]
        if kind == "heading":
            doc.add_paragraph(block[1], style=self._safe_style(doc, "Heading 1"))
        elif kind == "list":
            for n, item in enumerate(block[1], 1):
                doc.add_paragraph(f"{n}. {item}" if block[2] else f"\u2022 {item}")
        elif kind == "table":
            headings, rows = block[1], block[2]
            for row in rows:
                doc.add_paragraph(" - ".join(
                    f"{h}: {c}" if h.strip() else c
                    for h, c in zip(headings, row) if c.strip()
                ))
        elif kind == "code":
            doc.add_paragraph("\n".join(_wrap_lines(block[1], 96)),
                              style=self._safe_style(doc, "Code Block"))
        else:
            doc.add_paragraph(block[1])

    # -- floor: no styles, no tables, nothing that needs raw OOXML ----------

    def _plain_docx(self) -> bytes:
        doc = DocxFile()
        if self.document_title:
            doc.add_paragraph(self.document_title)
        for block in self.blocks:
            kind = block[0]
            if kind == "table":
                headings, rows = block[1], block[2]
                for row in rows:
                    for h, c in zip(headings, row):
                        if c.strip():
                            doc.add_paragraph(f"{h}: {c}" if h.strip() else c)
            elif kind == "list":
                for n, item in enumerate(block[1], 1):
                    doc.add_paragraph(f"{n}. {item}" if block[2] else f"- {item}")
            else:
                doc.add_paragraph(block[1])
        buf = BytesIO()
        doc.save(buf)
        return buf.getvalue()

# Excel reads a cell as a formula when it opens with one of these, and a CSV is
# the classic way that gets in: the document says =cmd|'/c calc'!A0 because a
# model put it there, and the reader's spreadsheet offers to run it. Leading
# whitespace and tabs are stripped before the check because Excel strips them too.
_FORMULA_LEAD = ("=", "+", "-", "@")

# ...but a negative number legitimately starts with "-", and quoting those as
# text would break every numeric column. Currency and grouping marks come off
# first so "-$1,204.00" still reads as a number.
_NUMERIC_NOISE = str.maketrans("", "", "$\u20ac\u00a3\u00a5%, \u00a0_")


def _is_number(cell: str) -> bool:
    try:
        float(cell.translate(_NUMERIC_NOISE))
        return True
    except ValueError:
        return False


def _defuse(cell: str) -> str:
    lead = cell.lstrip()[:1]
    if lead in _FORMULA_LEAD and not (lead in ("+", "-") and _is_number(cell)):
        # The apostrophe is the standard mitigation and the one Excel, Sheets
        # and LibreOffice all honour. The cost is that it stays visible: a CSV
        # import keeps it in the cell rather than consuming it as a text marker,
        # so =SUM(C2:C4) written by a model reaches the reader as '=SUM(C2:C4).
        return "'" + cell
    return cell


def _flatten(text: str) -> str:
    """Inline Markdown as something worth reading in a spreadsheet cell.

    The prose renderers can leave **bold** alone because their output is prose.
    A cell is data, and `**Payments**` in a column of team names is just noise.
    """
    out = []
    for chunk, _bold, _italic, _mono, href in _segments(text):
        out.append(f"{chunk} ({href})" if href and href != chunk else chunk)
    return "".join(out)


class CsvRenderer(DocumentRenderer):
    """The document's first table, as a spreadsheet.

    Subclass to retune: `delimiter = "\\t"` gives a TSV renderer.
    """

    delimiter = ","

    # Excel decodes a BOM-less UTF-8 CSV as the local ANSI codepage, which turns
    # every non-ASCII name in the table into mojibake. Off if the consumer is a
    # parser rather than a person.
    bom = True

    # Off if the tables are trusted and you would rather keep =SUM(B2:B9) intact.
    defuse_formulas = True

    # -- collection ---------------------------------------------------------

    def reset(self) -> None:
        self.headings: list[str] = []
        self.rows: list[list[str]] = []
        self.tables = 0

    def heading(self, text, level) -> None:
        pass  # a CSV has no room for anything but the table

    def paragraph(self, text) -> None:
        pass

    def item_list(self, items, ordered) -> None:
        pass

    def code_block(self, code, language) -> None:
        pass

    def table(self, headings, rows) -> None:
        headings, rows = _table_data(headings, rows)
        if not rows and not any(h.strip() for h in headings):
            # _table_data() turns table(None, None) into one blank column, and a
            # model emits those. Taking it as "the first table" would hand back
            # an empty file while a real table sat further down the document.
            return
        self.tables += 1
        if self.tables == 1:
            self.headings, self.rows = headings, rows

    # -- output -------------------------------------------------------------

    def finish(self) -> bytes:
        if not self.tables:
            logger.warning("%r has no table; the CSV is empty", self.document_title)
            return b""
        if self.tables > 1:
            logger.info("%r has %d tables; the CSV carries the first",
                        self.document_title, self.tables)

        buf = StringIO(newline="")
        writer = csv.writer(buf, delimiter=self.delimiter,
                            quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
        if any(h.strip() for h in self.headings):
            # All-blank headings mean the model gave none, and a leading empty
            # row would become the column names in whatever reads this.
            writer.writerow([self._cell(h) for h in self.headings])
        writer.writerows([self._cell(c) for c in row] for row in self.rows)

        text = "\ufeff" + buf.getvalue() if self.bom else buf.getvalue()
        # "replace", not strict: a lone surrogate survives json.loads() of model
        # output and would raise here, the way TxtRenderer.finish() still can.
        return text.encode("utf-8", "replace")

    def _cell(self, value: str) -> str:
        """One model-authored cell as one spreadsheet cell."""
        text = _flatten(value)
        # Cell-internal newlines are legal and csv.writer quotes them, but a
        # bare \r inside a \r\n-terminated file confuses line-splitting readers.
        text = re.sub(r"\r\n?", "\n", text)
        return _defuse(text) if self.defuse_formulas else text